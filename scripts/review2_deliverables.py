from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parent.parent
SUBMISSION = ROOT / 'submission_artifacts'
VERIFY = ROOT / 'verification_evidence'


def set_doc_title(doc, title: str):
    doc.add_heading(title, 0)
    doc.add_paragraph('QuakeWatch Review 2 – evidence package compiled from the project source and the available verification logs.')


def compile_summary_lines():
    summary = [
        'Project: QuakeWatch',
        'Technology stack: Django + Python + ArangoDB + Bootstrap + HTML/CSS/JS',
        'Database target: ArangoDB; project docs and schema logs identify quakewatch as the primary local database.',
        'Evidence sources: README, docs/ARCHITECTURE.md, docs/AQL_QUERIES.md, verification_evidence/queries/baseline_2026-10-07_aql-catalog-run.txt, verification_evidence/tests/baseline_2026-10-07_phase0-report.md, verification_evidence/database/environment_record.md.',
        'Result: The project contains a real ArangoDB-backed design and more than 15 catalogued complex AQL queries. The current Windows environment does not expose a live running ArangoDB instance, so live database counts, transactions, and export generation remain blocked until the database is started in a verified isolated environment.',
    ]
    return summary


def generate_review_docx():
    out = SUBMISSION / 'Review2_QuakeWatch.docx'
    doc = Document()
    set_doc_title(doc, 'Review 2 – QuakeWatch')

    for p in compile_summary_lines():
        doc.add_paragraph(p)

    doc.add_heading('1. Project brief', level=1)
    text = (
        'QuakeWatch is a Django application that ingests USGS earthquake data, normalizes and deduplicates it, stores revision-aware records in ArangoDB, and shows them through a dashboard, map, analytics, alerts, research, and operations interfaces. '
        'The project is explicitly documented as an ArangoDB-only stateful service with Django as the application layer.'
    )
    doc.add_paragraph(text)

    doc.add_heading('2. Database implementation status', level=1)
    rows = [
        ['Item', 'Finding', 'Evidence'],
        ['Database', 'ArangoDB is the project’s sole stateful database and is used for domain data, accounts, sessions, cache, metrics, jobs and scheduler state.', 'README.md; config/env.py; docs/ARCHITECTURE.md'],
        ['Schema and migration', 'Schema, indexes, view creation and graph setup are handled by project migration tooling and documented in the query baseline.', 'verification_evidence/queries/baseline_2026-10-07_aql-catalog-run.txt'],
        ['Live DB status', 'No live ArangoDB listener was reachable in the current Windows environment, so record counts and export generation remain blocked until an isolated dev DB is running.', 'verification_evidence/database/environment_record.md'],
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    header = table.rows[0].cells
    for j, item in enumerate(rows[0]):
        header[j].text = item
    for row in rows[1:]:
        cells = table.add_row().cells
        for j, item in enumerate(row):
            cells[j].text = item

    doc.add_heading('3. Record count and evidence', level=1)
    doc.add_paragraph('The project baseline report confirms the benchmark performed synthetic dataset seeding at 500, 10,000 and 100,000 events. The current environment did not expose a running database, so the 500-record requirement is not currently verified against a live dev database.')
    doc.add_paragraph('Historical evidence: the baseline benchmark log records “== 500 events ==”, “seeded 500 events in 0.2s”, and the benchmark output confirms the indexed query patterns were exercised.')

    doc.add_heading('4. Indexes, views and AQL coverage', level=1)
    doc.add_paragraph('The project defines an ArangoDB schema with persistent indexes, TTL indexes, edge collections and a search view. The AQL catalogue lists 53 queries and includes 15 selected high-value queries suitable for the Review 2 presentation. The baseline run reports 47 query entries and identifies the index choices used in the optimizer for the real data path.')

    qfile = ROOT / 'submission_artifacts' / '06_aql_queries' / 'selected_15_queries.md'
    qtext = qfile.read_text(encoding='utf-8') if qfile.exists() else 'Selected queries not available.'
    doc.add_paragraph(qtext[:3500])

    doc.add_heading('5. CRUD and backend connectivity', level=1)
    doc.add_paragraph('The project code includes Django CRUD surfaces for alert rules, saved searches, research snapshots and operations. The current working copy has code-level evidence for create/read/update/delete paths, but live CRUD execution remains blocked because the target database is not running and no ArangoDB endpoint is available in the current environment.')

    doc.add_heading('6. Testing and verification', level=1)
    doc.add_paragraph('Historical baseline results: 601 passed, 0 failed, 0 errors, coverage 93.96%. Browser/accessibility: 58 passed. The real environment however was a separate qwb baseline, not the current local Windows workspace, so these results must be read as historical evidence rather than active current runtime validation.')

    doc.add_heading('7. Compliance summary', level=1)
    doc.add_paragraph('This package satisfies the source-analysis and evidence-documentation component of Review 2 and clearly identifies blocked live validation items: current DB counts, DB export, and transaction-level execution. The project does contain the required broad design, indexes, AQL catalogue, and test evidence, but the current environment is not safe to claim end-to-end live verification.')

    doc.save(out)
    return out


def generate_query_summary_pdf():
    output = SUBMISSION / '6_Queries.pdf'
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('Title', parent=styles['Title'], fontName='Helvetica-Bold', fontSize=18, leading=22)
    body_style = ParagraphStyle('Body', parent=styles['BodyText'], fontName='Helvetica', fontSize=10, leading=14)
    story = [Paragraph('QuakeWatch – 15 AQL Query Summary', title_style), Spacer(1, 12)]

    query_names = [
        'Q01 analytics_avg_magnitude_by_region', 'Q02 analytics_dashboard_summary', 'Q03 analytics_network_stats',
        'Q04 analytics_magnitude_changes', 'Q05 analytics_ingestion_latency', 'Q06 analytics_significant_by_week',
        'Q07 alerts_user_rules', 'Q08 alerts_matched_events_graph', 'Q09 event_graph_relationships',
        'Q10 events_in_viewport', 'Q11 events_within_radius', 'Q12 events_location_text', 'Q13 ops_run_stats',
        'Q14 analytics_depth_histogram', 'Q15 event_nearby_space_time'
    ]
    for idx, name in enumerate(query_names, start=1):
        story.append(Paragraph(f'{idx}. {name}', body_style))
        story.append(Paragraph('Purpose: complex AQL routine using aggregation, geo search, graph traversal, or revision analysis.', body_style))
        story.append(Spacer(1, 6))

    story.append(Paragraph('Evidence source: docs/AQL_QUERIES.md and verification_evidence/queries/baseline_2026-10-07_aql-catalog-run.txt', body_style))
    story.append(Paragraph('Status: the catalogue and query logs confirm these query families exist; active DB execution is blocked in the current environment until the ArangoDB service is started in an isolated environment.', body_style))
    doc = SimpleDocTemplate(str(output), pagesize=A4)
    doc.build(story)
    return output


def generate_review_pdf():
    output = SUBMISSION / 'Review2_QuakeWatch.pdf'
    styles = getSampleStyleSheet()
    body = ParagraphStyle('Body', parent=styles['BodyText'], fontName='Helvetica', fontSize=10, leading=14)
    title = ParagraphStyle('Title', parent=styles['Title'], fontName='Helvetica-Bold', fontSize=18, leading=22)
    story = [
        Paragraph('QuakeWatch Review 2', title),
        Spacer(1, 12),
        Paragraph('This report is based on the real QuakeWatch source tree, the project documentation, and verification logs already present in the workspace. It explicitly distinguishes confirmed evidence from blocked live validation that requires an active ArangoDB instance.', body),
        Spacer(1, 12),
        Paragraph('Project: QuakeWatch | Stack: Django, Python, ArangoDB, HTML/CSS/JS, Bootstrap | Source: README.md, docs/ARCHITECTURE.md, docs/AQL_QUERIES.md, submission_artifacts/06_aql_queries/selected_15_queries.md.', body),
        Spacer(1, 12),
        Paragraph('Database evidence: the project is deliberately designed around ArangoDB as the sole stateful database. The baseline AQL-catalogue log confirms schema creation, index creation, view creation, and query optimizer choices for 47 named queries. The current Windows workspace however does not expose a live ArangoDB service, so record counts, transaction execution, and export creation remain blocked and must be re-run against a verified local database.', body),
        Spacer(1, 12),
        Paragraph('Query evidence: multiple query families exist for alerts, analytics, event lookups, geo search, revision inspection, and operations. The catalogue includes 53 named queries; the project also has a selected 15-query set ready for Review 2 presentation. Historical output shows query execution status and index usage across a seeded benchmark database.', body),
        Spacer(1, 12),
        Paragraph('Testing evidence: Phase 0 baseline results recorded 601 passed, 0 failed, 0 errors, and 93.96% coverage. Browser/accessibility checks also passed in the historical baseline. These results are valid as historical evidence but not as a current live verification claim for the present working copy, which does not presently have a running ArangoDB listener.', body),
        Spacer(1, 12),
        Paragraph('Outcome: the project is validly documented and the documentation/evidence set is sufficient for Review 2 source analysis; the remaining live-environment steps are blocked by the absence of a reachable ArangoDB service and must be executed in a separate, safely managed local environment.', body),
    ]
    doc = SimpleDocTemplate(str(output), pagesize=A4)
    doc.build(story)
    return output


def generate_testcases_xlsx():
    output = SUBMISSION / '7_TestCases.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = 'Test Cases'
    headers = ['Test ID', 'Requirement', 'Status', 'Evidence', 'Notes']
    ws.append(headers)
    rows = [
        ['TC-01', 'Committed transaction writes all touched collections', 'Executed / historical valid', 'project transaction docs; baseline logs', 'Requires live DB to re-execute.'],
        ['TC-02', 'Failed transaction leaves no partial writes', 'Executed / historical valid', 'project transaction docs; baseline logs', 'Currently blocked in the active workspace.'],
        ['TC-03', 'Concurrent writers do not lose updates', 'Proposed', 'project test strategy', 'Not re-run in the current environment.'],
        ['TC-04', 'Invalid ingest record is dead-lettered', 'Proposed', 'ingestion design and dead_letters collection', 'Dead-letter path exists in code.'],
        ['TC-05', 'User can log in', 'Proposed', 'Django auth and Arango accounts', 'Needs active DB.'],
        ['TC-06', 'User can create/read/update/delete alert rule', 'Proposed / code-level implemented', 'apps/alerts/api.py', 'Live CRUD execution blocked.'],
        ['TC-07', 'User can save/list/run/delete search', 'Proposed / code-level implemented', 'apps/earthquakes/api.py', 'Live CRUD execution blocked.'],
        ['TC-08', 'Research export has manifest and checksum', 'Proposed', 'apps/research', 'Needs live export run.'],
        ['TC-09', 'Map loads and shows events', 'Proposed', 'map and API endpoints exist', 'Browser run blocked.'],
        ['TC-10', 'Public pages no console/CSP errors', 'Historical pass', 'baseline browser logs', '58 passed in historical suite.'],
        ['TC-11', 'Dashboard summary works', 'Historical pass', 'baseline tests / AQL catalogue', 'Historical evidence only.'],
        ['TC-12', 'Analytics queries run correctly', 'Historical pass', 'catalogue and query logs', 'Values are not current live results.'],
        ['TC-13', 'Index use matches expectation', 'Historical pass', 'aql_catalog baseline log', 'Two queries were flagged on empty DB as data-dependent, not a general failure.'],
        ['TC-14', 'Query plan remains safe for filter conditions', 'Executed / historical valid', 'AQL catalog baseline', 'Query plans captured from log.'],
        ['TC-15', 'Invalid events rejected by validator', 'Proposed', 'schema validation in Arango collections', 'Needs live DB validation.'],
        ['TC-16', 'Duplicate source identity is rejected', 'Proposed', 'unique index and schema rules', 'Current environment blocked.'],
        ['TC-17', 'Non-staff users cannot call admin endpoints', 'Proposed', 'Authorization rules in code', 'Needs live auth runtime.'],
        ['TC-18', '500 record minimum requirement is satisfied', 'Blocked', 'synthetic benchmark only', 'No live DB currently reachable.'],
        ['TC-19', '15+ complex AQL queries are available', 'Historical pass', 'docs/AQL_QUERIES.md and selected_15_queries.md', 'Confirmed by project documentation.'],
        ['TC-20', 'Database export is generated', 'Blocked', 'no live DB endpoint', 'Requires identified ArangoDB environment.'],
    ]
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.font = cell.font.copy(bold=True)
    ws.freeze_panes = 'A2'
    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 38
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 42
    ws.column_dimensions['E'].width = 30
    wb.save(output)
    return output


def write_supporting_summary():
    txt = '''QuakeWatch Review 2 – evidence status

This file consolidates the evidence already present in the repository and provides the current status of live verification.

Confirmed from source and logs:
- ArangoDB is the sole stateful back-end.
- Project includes a named AQL catalogue and 53 query definitions.
- A selected 15-query set is prepared in submission_artifacts/06_aql_queries/selected_15_queries.md.
- Historical benchmark/test baseline recorded 601 passed, 0 failed, coverage 93.96%.
- Browser/accessibility baseline recorded 58 passed.

Blocked in the current workspace:
- No reachable local ArangoDB listener on 127.0.0.1:8529.
- No active Docker Desktop or qwb DB environment is running.
- Database counts, live CRUD verification, transaction execution, and DB export creation require a verified isolated database environment before safe runtime claims can be made.

The generated Review 2 deliverables therefore include real evidence from the repository while accurately marking unverified live tasks as blocked.
'''
    path = VERIFY / 'review2_status_summary.md'
    path.write_text(txt, encoding='utf-8')
    return path


def main():
    SUBMISSION.mkdir(exist_ok=True)
    VERIFY.mkdir(exist_ok=True)
    generate_review_docx()
    generate_review_pdf()
    generate_query_summary_pdf()
    generate_testcases_xlsx()
    write_supporting_summary()
    print('Generated:')
    for p in [
        SUBMISSION / 'Review2_QuakeWatch.docx',
        SUBMISSION / 'Review2_QuakeWatch.pdf',
        SUBMISSION / '6_Queries.pdf',
        SUBMISSION / '7_TestCases.xlsx',
        VERIFY / 'review2_status_summary.md',
    ]:
        print(p)


if __name__ == '__main__':
    main()
